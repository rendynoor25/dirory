#!/usr/bin/env python
"""
Read the Kalua Studio vendor sheet and emit the contact data the outreach builder
needs: proposals/outreach/contacts.json

Why a separate extractor rather than typing the numbers into the builder: there are
thirty of them, they arrive in five different formats (`62 878-7480-0700`,
`6281110119278.0`, `86 188 3291 7521`) and one wrong digit sends a proposal to a
stranger. Reading them mechanically removes that class of mistake.

Only rows that map to a slug are emitted — the sheet also contains permit
consultants, structural analysts, web developers, an interior contractor, two
Autodesk/SketchUp resellers and a signage fabricator, none of which is a Dirory
vendor.

Usage:  python scripts/extract-kalua-contacts.py
"""

from __future__ import annotations

import json
import os
import re
import sys

import openpyxl

SHEET = r"C:\Users\HP245\Downloads\Database Vendor (by Kalua Studio).xlsx"
OUT = r"D:\Dirory\dirory\proposals\outreach\contacts.json"

# company name in the sheet -> slug used by build-outreach-proposals.mjs
SLUGS = {
    "PT SURYA UTAMA FIBERTEK": "surya-utama-fibertek",
    "Waveboard": "waveboard",
    "Duma": "duma",
    "Immortal Marble & Granite": "immortal-marble-granite",
    "Panelku": "panelku",
    "Muda Mandiri Perkasa": "muda-mandiri-perkasa",
    "Putra Glass Bandung": "putra-glass-bandung",
    "Modico": "modico",
    "Greenpro Indonesia": "greenpro-indonesia",
    "ReHolz": "reholz",
    "Astral Aluminium": "astral-aluminium",
    "PT Onna Prima Utama": "onna-prima-utama",
    "Atipa": "atipa",
    "Skayu International": "skayu-international",
    "Alubeen": "alubeen",
    "Artgrass Asia Rumput Sintetis": "artgrass-asia",
    "Roman": "roman",
    "Distributor Epoxy 789": "distributor-epoxy-789",
    "Moire Rugs": "moire-rugs",
    "Bottega & Artisan (PT Berkat Artisan Indonesia)": "bottega-artisan",
    "Bellaquattro": "bellaquattro",
    "Higold Indonesia": "higold-indonesia",
    "Daeji Vinyl": "daeji-vinyl",
    "PT Altara Indonesia Sukses": "altara-indonesia",
    "Berdikari Tunggal Perkasa": "berdikari-tunggal-perkasa",
    "Paranti Sofa": "paranti-sofa",
    "Sanwoo Electronics": "sanwoo-electronics",
    "Keding HPL Bandung": "keding-hpl-bandung",
    "Lalapis": "lalapis",
    "Solusi Atap Prima": "solusi-atap-prima",
}


def digits(raw) -> str:
    """Normalise a phone number to bare digits, suitable for a wa.me link.

    The sheet mixes formats and Excel has turned several of them into floats, so
    `6281110119278.0` must not become `62811101192780`.
    """
    s = str(raw or "").strip()
    s = re.sub(r"\.0+$", "", s)          # drop Excel's float tail
    s = re.sub(r"[^\d]", "", s)          # spaces, dashes, plus
    return s


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    ws = openpyxl.load_workbook(SHEET, data_only=True).worksheets[0]
    out: dict[str, dict] = {}
    unmatched: list[str] = []

    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        cells = ["" if c is None else str(c).strip() for c in row]
        # Column layout in this sheet is  A=?, B=No, C=Nama Vendor, D=PIC,
        # E=Nomor HP, F=Catalog, G=Deskripsi Barang — so the name is index 2.
        cell = lambda i: cells[i] if len(cells) > i else ""
        nama = re.sub(r"\s+", " ", cell(2))
        pic = cell(3)
        hp = cell(4)
        cat = cell(5)
        if not nama:
            continue
        slug = SLUGS.get(nama)
        if not slug:
            if nama.lower() not in ("nama vendor",):
                unmatched.append(nama)
            continue
        out[slug] = {
            "company": nama,
            "pic": pic,
            "phone": digits(hp),
            "sheet_note": cat if cat else "",
        }

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)

    print("wrote %s" % OUT)
    print("  matched: %d vendors" % len(out))
    missing = [s for s in SLUGS.values() if s not in out]
    if missing:
        print("  !! slugs with no row in the sheet: %s" % ", ".join(missing))
    print()
    for slug, v in sorted(out.items()):
        flag = "" if v["phone"].startswith("62") else "   <-- not an Indonesian number"
        print("  %-30s %-14s %-16s%s" % (slug, v["pic"][:14], v["phone"], flag))
    if unmatched:
        print("\n  rows not mapped (expected — not Dirory vendors):")
        for n in unmatched:
            print("    - %s" % n)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
