"""Recalculate every formula in a Dirory RAB workbook and compare it with the
value the writer cached. Supports the formula subset rab-xlsx.js writes:
cell and cross-sheet references, + - * /, SUM(range), IF(a=0,b,c) and
ROUNDDOWN(x,n).   Usage: python scripts/check-xlsx.py book.xlsx"""
import math, re, sys
import openpyxl

path = sys.argv[1]
wf = openpyxl.load_workbook(path)
wv = openpyxl.load_workbook(path, data_only=True)
memo = {}

def value(sheet, coord):
    key = (sheet, coord)
    if key in memo:
        return memo[key]
    cell = wf[sheet][coord]
    v = cell.value
    if isinstance(v, str) and v.startswith("="):
        v = evaluate(sheet, v[1:])
    memo[key] = 0 if v is None else v
    return memo[key]

def col_range(sheet, a, b):
    ws = wf[sheet]
    total = 0
    for row in ws[a.replace("$", ""):b.replace("$", "")]:
        for c in row:
            x = value(sheet, c.coordinate)
            total += x if isinstance(x, (int, float)) else 0
    return total

def evaluate(sheet, f):
    expr = f
    # SUM(A1:B2)
    expr = re.sub(r"SUM\(([A-Z$]+\d+):([A-Z$]+\d+)\)", lambda m: repr(col_range(sheet, m.group(1), m.group(2))), expr)
    expr = expr.replace("ROUNDDOWN", "_rounddown").replace("IF(", "_if(")
    expr = re.sub(r"(?:'([^']+)'!|([A-Za-z_][A-Za-z0-9_]*)!)?\$?([A-Z]{1,3})\$?(\d+)(?![\w(])",
                  lambda m: f"_v({(m.group(1) or m.group(2) or sheet)!r},{(m.group(3) + m.group(4))!r})", expr)
    expr = expr.replace("=", "==").replace("<==", "<=").replace(">==", ">=")
    return eval(expr, {"_v": value, "_if": lambda c, a, b: a if c else b,
                       "_rounddown": lambda x, n: math.floor(x * 10 ** n) / 10 ** n})

bad = checked = 0
for ws in wf.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                got = value(ws.title, c.coordinate)
                cached = wv[ws.title][c.coordinate].value
                checked += 1
                if not isinstance(got, (int, float)) or cached is None or abs(got - cached) > max(1, abs(cached) * 1e-9):
                    bad += 1
                    if bad <= 10:
                        print(f"MISMATCH {ws.title}!{c.coordinate} {c.value} -> {got} cached {cached}")
rekap = wf["REKAP"]
for row in rekap.iter_rows():
    if row[1].value in ("JUMLAH", "TOTAL TERMASUK PPN"):
        print(f"{row[1].value}: {value('REKAP', row[2].coordinate):,.0f}")
print(f"{checked} formulas recalculated, {bad} mismatches")
sys.exit(1 if bad else 0)
